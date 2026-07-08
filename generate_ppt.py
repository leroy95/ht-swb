from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor

OUT = r"D:\\Work\\Notes\\Personal_Notes\\手工文创\\手作文创项目概览.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

COLORS = {
    "navy": RGBColor(0x2F, 0x3C, 0x7E),
    "coral": RGBColor(0xF9, 0x61, 0x67),
    "gold": RGBColor(0xF9, 0xE7, 0x95),
    "cream": RGBColor(0xFC, 0xF6, 0xF5),
    "sand": RGBColor(0xF4, 0xEC, 0xE8),
    "ink": RGBColor(0x2B, 0x2B, 0x2B),
    "muted": RGBColor(0x6B, 0x72, 0x80),
    "rose": RGBColor(0xFF, 0xE7, 0xE1),
    "sage": RGBColor(0xA7, 0xBE, 0xAE),
    "white": RGBColor(0xFF, 0xFF, 0xFF),
    "line": RGBColor(0xE5, 0xD6, 0xD1),
}

TITLE_FONT = "Microsoft YaHei"
BODY_FONT = "Microsoft YaHei"


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(slide, text, x, y, w, h, size=18, color=None, bold=False,
             font=BODY_FONT, align=PP_ALIGN.LEFT, italic=False, name=None,
             margin_left=0.08, margin_right=0.08, margin_top=0.05, margin_bottom=0.05):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin_left)
    tf.margin_right = Inches(margin_right)
    tf.margin_top = Inches(margin_top)
    tf.margin_bottom = Inches(margin_bottom)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color or COLORS["ink"]
    if name:
        tb.name = name
    return tb


def add_bullets(slide, items, x, y, w, h, size=17, color=None, level0_indent=0.22):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.font.name = BODY_FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color or COLORS["ink"]
        p.space_after = Pt(6)
        p.line_spacing = 1.15
        p.left_margin = Inches(level0_indent)
        p.first_line_indent = Inches(-0.14)
    return tb


def add_round_rect(slide, x, y, w, h, fill, line=None, radius=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line or fill
    shp.line.width = Pt(1)
    return shp


def add_circle(slide, x, y, d, fill, text=None, text_size=18, text_color=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = fill
    if text:
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text
        run.font.name = TITLE_FONT
        run.font.size = Pt(text_size)
        run.font.bold = True
        run.font.color.rgb = text_color or COLORS["white"]
    return shp


def add_metric(slide, x, y, w, h, number, label, fill, number_color=COLORS["navy"]):
    add_round_rect(slide, x, y, w, h, fill, line=fill)
    add_text(slide, number, x + 0.18, y + 0.18, w - 0.36, 0.60, size=28, color=number_color, bold=True, font=TITLE_FONT)
    add_text(slide, label, x + 0.18, y + 0.82, w - 0.36, h - 0.92, size=13, color=COLORS["muted"])


def add_footer(slide, text, dark=False):
    color = COLORS["cream"] if dark else COLORS["muted"]
    add_text(slide, text, 0.65, 7.0, 12.0, 0.28, size=9, color=color)


# Slide 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, COLORS["navy"])
add_round_rect(slide, 0.65, 0.55, 3.1, 0.42, COLORS["coral"], line=COLORS["coral"], radius=False)
add_text(slide, "广胜手作计划", 0.82, 1.15, 5.5, 0.7, size=28, color=COLORS["gold"], bold=True, font=TITLE_FONT)
add_text(slide, "手作文创项目概览", 0.82, 1.78, 7.8, 1.0, size=30, color=COLORS["white"], bold=True, font=TITLE_FONT)
add_text(slide, "围绕“产品库—计划库—执行手册”三层结构，搭建双线文创业务：\n一条做洪洞/广胜寺地域文创，一条做手工高端与体验型产品。", 0.82, 2.75, 6.4, 1.4, size=18, color=COLORS["cream"])
for x, y, d, c in [(9.65, 1.08, 1.35, COLORS["coral"]), (10.78, 2.12, 1.02, COLORS["gold"]), (8.95, 3.25, 1.72, COLORS["sage"]), (10.45, 4.18, 1.25, COLORS["rose"]), (11.52, 3.38, 0.72, COLORS["cream"])]:
    add_circle(slide, x, y, d, c)
add_round_rect(slide, 8.1, 1.45, 3.85, 4.35, COLORS["cream"], line=COLORS["cream"])
add_text(slide, "项目关键信号", 8.45, 1.82, 2.8, 0.45, size=20, color=COLORS["navy"], bold=True, font=TITLE_FONT)
add_metric(slide, 8.45, 2.42, 1.55, 1.35, "5", "核心手工产品线", COLORS["rose"])
add_metric(slide, 10.18, 2.42, 1.55, 1.35, "4", "地域文创方向", COLORS["gold"])
add_metric(slide, 8.45, 3.96, 1.55, 1.35, "2", "加速专题：DIY / 唐装合作", COLORS["sand"])
add_metric(slide, 10.18, 3.96, 1.55, 1.35, "3", "推进层：产品 / 计划 / 执行", COLORS["rose"])
add_footer(slide, "来源：README、业务逻辑总览、产品索引、计划索引、制作落地执行方案", dark=True)

# Slide 2
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, COLORS["cream"])
add_text(slide, "01｜业务结构：先定义做什么，再决定怎么推", 0.72, 0.58, 8.4, 0.6, size=28, color=COLORS["navy"], bold=True, font=TITLE_FONT)
add_text(slide, "文档体系已经把业务判断收敛到统一入口，便于后续产品、合作与执行动作按同一逻辑推进。", 0.72, 1.1, 8.7, 0.42, size=15, color=COLORS["muted"])
cols = [
    (0.82, "产品库", COLORS["rose"], ["回答“做什么”", "围绕用户、价格带、差异化与版本演进", "覆盖核心产品线与地域文创矩阵"]),
    (4.42, "计划库", COLORS["gold"], ["回答“怎么推进”", "通过目标、范围、选型、风险控制推进", "只引用版本，不重复产品正文"]),
    (8.02, "执行手册", COLORS["sand"], ["回答“怎么落地”", "聚焦打样、试产、质检、入库与异常处理", "强调能交接、能追责、能回退"]),
]
for x, title, fill, bullets in cols:
    add_round_rect(slide, x, 1.75, 3.0, 4.6, fill, line=fill)
    add_circle(slide, x + 0.22, 1.98, 0.48, COLORS["navy"], text="●", text_size=12, text_color=COLORS["gold"])
    add_text(slide, title, x + 0.84, 1.98, 1.8, 0.34, size=21, color=COLORS["navy"], bold=True, font=TITLE_FONT)
    add_bullets(slide, bullets, x + 0.26, 2.62, 2.42, 2.2, size=15)
add_text(slide, "统一判断点", 0.95, 6.72, 1.35, 0.3, size=13, color=COLORS["coral"], bold=True)
labels = ["用户", "价格", "渠道", "供应链", "版本"]
for i, lab in enumerate(labels):
    px = 2.12 + i * 2.0
    add_circle(slide, px, 6.48, 0.44, COLORS["navy"], text=lab[0], text_size=14)
    add_text(slide, lab, px + 0.52, 6.56, 0.9, 0.22, size=12, color=COLORS["ink"])
add_footer(slide, "判断规则摘自：设计/00_索引导航/业务逻辑总览.md")

# Slide 3
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, COLORS["cream"])
add_text(slide, "02｜产品矩阵：一条做地域引流，一条做手工高端", 0.72, 0.58, 8.7, 0.6, size=28, color=COLORS["navy"], bold=True, font=TITLE_FONT)
add_text(slide, "核心产品线负责品牌高度，地域文创负责可传播、可补货、可快速验证的商品体系。", 0.72, 1.1, 8.6, 0.36, size=15, color=COLORS["muted"])
add_round_rect(slide, 0.82, 1.75, 5.9, 4.95, COLORS["rose"], line=COLORS["rose"])
add_text(slide, "手工高端线", 1.08, 2.0, 2.1, 0.4, size=22, color=COLORS["navy"], bold=True, font=TITLE_FONT)
manual_items = [
    "法式刺绣高定饰品",
    "天然植物染真丝丝巾",
    "浮雕艺术香薰蜡烛",
    "微型钩编植物艺术首饰",
    "环氧树脂干花艺术家居",
]
for idx, item in enumerate(manual_items):
    y = 2.62 + idx * 0.66
    add_circle(slide, 1.12, y, 0.34, COLORS["coral"], text=str(idx + 1), text_size=12)
    add_text(slide, item, 1.56, y - 0.01, 3.95, 0.28, size=15, color=COLORS["ink"])
add_text(slide, "定位关键词：高客单、礼赠、收藏、品牌形象款", 1.12, 6.03, 4.85, 0.32, size=13, color=COLORS["muted"])

add_round_rect(slide, 6.95, 1.75, 5.56, 4.95, COLORS["gold"], line=COLORS["gold"])
add_text(slide, "地域文创线", 7.22, 2.0, 2.0, 0.4, size=22, color=COLORS["navy"], bold=True, font=TITLE_FONT)
regional_cards = [
    ("洪洞地域文创", "首轮优先做纸品布艺样板", 7.22, 2.58),
    ("广胜寺纸品布艺", "走量与图案传播", 9.95, 2.58),
    ("广胜寺琉璃系列", "偏形象与工艺表达", 7.22, 4.18),
    ("广胜寺香道禅修", "偏体验与精神消费", 9.95, 4.18),
]
for title, desc, x, y in regional_cards:
    add_round_rect(slide, x, y, 2.35, 1.25, COLORS["cream"], line=COLORS["cream"])
    add_text(slide, title, x + 0.14, y + 0.16, 2.05, 0.34, size=15, color=COLORS["navy"], bold=True)
    add_text(slide, desc, x + 0.14, y + 0.56, 2.05, 0.42, size=11.5, color=COLORS["muted"])
add_footer(slide, "来源：产品索引、洪洞地域文创概览、各核心产品概览")

# Slide 4
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, COLORS["navy"])
add_text(slide, "03｜优先落地的增长引擎", 0.75, 0.62, 5.4, 0.6, size=28, color=COLORS["white"], bold=True, font=TITLE_FONT)
add_text(slide, "在主产品线之外，文档中已经给出两个更容易带动传播和成交的专题入口。", 0.75, 1.14, 6.6, 0.35, size=15, color=COLORS["cream"])
# DIY card
add_round_rect(slide, 0.88, 1.88, 5.75, 4.9, COLORS["cream"], line=COLORS["cream"])
add_circle(slide, 1.18, 2.18, 0.56, COLORS["coral"], text="DIY", text_size=12)
add_text(slide, "DIY体验套装", 1.92, 2.16, 2.4, 0.35, size=22, color=COLORS["navy"], bold=True, font=TITLE_FONT)
add_text(slide, "卖的不是单个成品，而是一次可完成、可拍照、可分享的体验。", 1.18, 2.72, 4.95, 0.45, size=14, color=COLORS["ink"])
add_bullets(slide, [
    "首发建议：琉璃色块拼接、金石拓印、古法合香 3 款主打",
    "适用场景：景区体验 / 电商礼盒 / 亲子活动 / 私域复购",
    "业务价值：高毛利、易传播、可做补充包和系列化",
], 1.18, 3.34, 4.95, 2.05, size=14)
add_round_rect(slide, 1.18, 5.88, 4.55, 0.52, COLORS["gold"], line=COLORS["gold"])
add_text(slide, "核心评价标准：用户能完成、愿意拍照、愿意再买", 1.38, 5.98, 4.15, 0.18, size=13, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)
# Tang card
add_round_rect(slide, 6.95, 1.88, 5.5, 4.9, COLORS["sand"], line=COLORS["sand"])
add_circle(slide, 7.25, 2.18, 0.56, COLORS["navy"], text="店", text_size=14)
add_text(slide, "唐装合作专题", 7.98, 2.16, 2.6, 0.35, size=22, color=COLORS["navy"], bold=True, font=TITLE_FONT)
add_text(slide, "当前以联名、配搭、门店合作为主，不与核心产品线混写。", 7.25, 2.72, 4.6, 0.45, size=14, color=COLORS["ink"])
add_bullets(slide, [
    "合作逻辑：先样品，再小批，再联名",
    "门店动作：先试戴和成交，再考虑会员、礼盒与放量",
    "所有合作要落到：谁做、做什么、多少钱、多久做完",
], 7.25, 3.34, 4.65, 2.05, size=14)
add_round_rect(slide, 7.25, 5.88, 4.15, 0.52, COLORS["rose"], line=COLORS["rose"])
add_text(slide, "弱反馈时，先改配搭与话术，不急于扩 SKU", 7.45, 5.98, 3.75, 0.18, size=13, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)
add_footer(slide, "来源：DIY体验套装设计、唐装合作计划 README", dark=True)

# Slide 5
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, COLORS["cream"])
add_text(slide, "04｜首发选择：先快验证，再做高单价形象款", 0.72, 0.58, 8.3, 0.6, size=28, color=COLORS["navy"], bold=True, font=TITLE_FONT)
add_text(slide, "目录中的建议非常一致：先用低风险、易传播、可补货的品类跑通，再把高客单产品做成品牌锚点。", 0.72, 1.1, 9.4, 0.36, size=15, color=COLORS["muted"])
# left timeline
add_text(slide, "推荐首发节奏", 0.95, 1.82, 1.9, 0.34, size=21, color=COLORS["coral"], bold=True, font=TITLE_FONT)
steps = [
    ("01", "地域文创先行", "纸品布艺首发：明信片、书签、便签本、帆布袋"),
    ("02", "DIY同步试水", "用 3 款套装验证体验、内容传播与客单价"),
    ("03", "手工高端树形象", "先做胸针、发夹、耳环等样板，控制小批量"),
]
for i, (num, title, desc) in enumerate(steps):
    y = 2.35 + i * 1.45
    add_circle(slide, 1.0, y, 0.58, COLORS["navy"], text=num, text_size=14)
    add_text(slide, title, 1.72, y + 0.02, 2.0, 0.26, size=17, color=COLORS["navy"], bold=True)
    add_text(slide, desc, 1.72, y + 0.36, 3.35, 0.46, size=13, color=COLORS["ink"])
    if i < 2:
        conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.29), Inches(y + 0.58), Inches(1.29), Inches(y + 1.28))
        conn.line.color.rgb = COLORS["line"]
        conn.line.width = Pt(2)
# right quadrant
add_text(slide, "优先级判断矩阵", 6.32, 1.82, 2.0, 0.34, size=21, color=COLORS["coral"], bold=True, font=TITLE_FONT)
add_round_rect(slide, 6.3, 2.28, 5.8, 3.95, COLORS["white"], line=COLORS["line"])
# axes
hline = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(7.15), Inches(4.23), Inches(11.45), Inches(4.23))
hline.line.color.rgb = COLORS["line"]
hline.line.width = Pt(1.6)
vline = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(9.3), Inches(2.75), Inches(9.3), Inches(5.7))
vline.line.color.rgb = COLORS["line"]
vline.line.width = Pt(1.6)
add_text(slide, "传播速度", 8.55, 2.42, 1.4, 0.22, size=11, color=COLORS["muted"])
add_text(slide, "成本与复杂度", 6.52, 4.38, 1.6, 0.22, size=11, color=COLORS["muted"])
points = [
    ("纸品布艺", 10.22, 2.98, COLORS["coral"]),
    ("DIY套装", 9.58, 3.55, COLORS["gold"]),
    ("刺绣饰品", 8.0, 4.88, COLORS["navy"]),
    ("香薰蜡烛", 8.68, 4.62, COLORS["sage"]),
]
for label, x, y, color in points:
    add_circle(slide, x, y, 0.28, color)
    add_text(slide, label, x + 0.36, y - 0.03, 1.4, 0.24, size=11.5, color=COLORS["ink"])
add_footer(slide, "来源：洪洞地域文创概览、DIY体验套装设计、法式刺绣高定饰品概览")

# Slide 6
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, COLORS["cream"])
add_text(slide, "05｜DIY 套装的商业价值与首发结构", 0.72, 0.58, 8.0, 0.6, size=28, color=COLORS["navy"], bold=True, font=TITLE_FONT)
add_text(slide, "DIY 套装兼具体验、内容传播和复购空间，是最适合快速做出用户反馈的产品形态之一。", 0.72, 1.1, 9.2, 0.38, size=15, color=COLORS["muted"])
for num, title, desc, x, fill in [("01", "引流款", "低门槛、快完成、适合景区现场体验", 0.95, COLORS["rose"]), ("02", "主销款", "体验完整、视觉强、适合线上售卖", 4.35, COLORS["gold"]), ("03", "高阶款", "更有挑战与收藏感，面向核心用户", 7.75, COLORS["sand"])]:
    add_round_rect(slide, x, 1.92, 2.7, 1.65, fill, line=fill)
    add_circle(slide, x + 0.18, 2.12, 0.4, COLORS["navy"], text=num, text_size=10)
    add_text(slide, title, x + 0.7, 2.1, 1.2, 0.25, size=18, color=COLORS["navy"], bold=True)
    add_text(slide, desc, x + 0.18, 2.58, 2.24, 0.45, size=12.5, color=COLORS["ink"])
for title, desc, x in [("琉璃色块拼接", "景区版 ¥68–88｜线上版 ¥98–128", 0.95), ("金石拓印", "景区版 ¥88–118｜线上版 ¥128–158", 4.35), ("古法合香", "景区版 ¥128–158｜线上版 ¥168–198", 7.75)]:
    add_round_rect(slide, x, 4.25, 2.7, 1.55, COLORS["white"], line=COLORS["line"])
    add_text(slide, title, x + 0.16, 4.41, 2.1, 0.25, size=16, color=COLORS["navy"], bold=True)
    add_text(slide, desc, x + 0.16, 4.8, 2.28, 0.35, size=11.5, color=COLORS["ink"])
add_round_rect(slide, 11.0, 1.92, 1.5, 3.88, COLORS["navy"], line=COLORS["navy"])
add_text(slide, "业务价值", 11.12, 2.12, 1.05, 0.25, size=16, color=COLORS["gold"], bold=True, font=TITLE_FONT)
add_bullets(slide, ["开箱、制作、成品都能出内容", "适合做系列、难度阶梯和补充包", "材料包模式更容易做出较高毛利"], 11.08, 2.58, 1.2, 2.15, size=10.5, color=COLORS["cream"], level0_indent=0.16)
add_footer(slide, "来源：DIY体验套装设计.md")

# Slide 7
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, COLORS["cream"])
add_text(slide, "06｜IP 授权边界：先做原创二创，再决定是否升级授权", 0.72, 0.58, 10.0, 0.6, size=27, color=COLORS["navy"], bold=True, font=TITLE_FONT)
add_text(slide, "IP 策略不是让项目停下来，而是帮助产品在安全边界内起步，并为后续授权谈判留下空间。", 0.72, 1.1, 9.8, 0.36, size=15, color=COLORS["muted"])
add_round_rect(slide, 0.95, 1.78, 3.85, 3.75, COLORS["rose"], line=COLORS["rose"])
add_text(slide, "可直接做", 1.18, 2.02, 1.4, 0.28, size=20, color=COLORS["navy"], bold=True, font=TITLE_FONT)
add_bullets(slide, ["古建筑外观与公共文化元素", "提取飞虹塔色彩、轮廓、纹样做原创设计", "产品说明写“灵感来源于广胜寺 / 洪洞”"], 1.1, 2.46, 3.45, 1.7, size=13.5)
add_round_rect(slide, 4.95, 1.78, 3.85, 3.75, COLORS["gold"], line=COLORS["gold"])
add_text(slide, "谨慎处理", 5.18, 2.02, 1.6, 0.28, size=20, color=COLORS["navy"], bold=True, font=TITLE_FONT)
add_bullets(slide, ["先做原创插画与二次创作，不直接复制照片", "产品卖起来后，再主动谈年度授权", "边界不清时优先改文案、改图稿、改品类"], 5.1, 2.46, 3.45, 1.7, size=13.5)
add_round_rect(slide, 8.95, 1.78, 3.4, 3.75, COLORS["sand"], line=COLORS["sand"])
add_text(slide, "不要直接做", 9.18, 2.02, 1.7, 0.28, size=20, color=COLORS["navy"], bold=True, font=TITLE_FONT)
add_bullets(slide, ["景区官方 Logo 或注册商标", "“官方授权 / 联名”字样", "直接使用他人摄影作品印刷"], 9.08, 2.46, 3.05, 1.7, size=13.5)
add_round_rect(slide, 1.0, 5.95, 11.3, 0.5, COLORS["navy"], line=COLORS["navy"])
add_text(slide, "推荐路径：第 1–6 个月先按原创二创路线启动；有销量和曝光后，再考虑主动谈授权升级。", 1.16, 6.06, 10.95, 0.16, size=13, color=COLORS["cream"], bold=True, align=PP_ALIGN.CENTER)
add_footer(slide, "来源：01_IP授权策略.md")

# Slide 8
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, COLORS["navy"])
add_text(slide, "07｜启动预算：先用小钱试错，再用回款滚动", 0.78, 0.7, 7.2, 0.58, size=29, color=COLORS["white"], bold=True, font=TITLE_FONT)
add_text(slide, "当前预算假设为 ¥30,000，重点不是一次铺满，而是控制风险、分批下单。", 0.78, 1.28, 7.2, 0.34, size=16, color=COLORS["cream"])
add_round_rect(slide, 0.92, 1.9, 3.2, 4.55, COLORS["cream"], line=COLORS["cream"])
add_text(slide, "预算结构", 1.18, 2.12, 1.6, 0.28, size=21, color=COLORS["navy"], bold=True, font=TITLE_FONT)
add_metric(slide, 1.18, 2.58, 1.15, 1.08, "59%", "洪洞线样品与首批备货", COLORS["rose"])
add_metric(slide, 2.52, 2.58, 1.15, 1.08, "10%", "设计外包", COLORS["gold"])
add_metric(slide, 1.18, 3.86, 1.15, 1.08, "10%", "包装 + 拍摄", COLORS["sand"])
add_metric(slide, 2.52, 3.86, 1.15, 1.08, "10%", "创始人生存费", COLORS["rose"])
add_text(slide, "其余预算用于手工线材料、营销和机动储备。", 1.18, 5.23, 2.4, 0.34, size=12.5, color=COLORS["muted"])
add_round_rect(slide, 4.45, 1.9, 4.0, 4.55, COLORS["sand"], line=COLORS["sand"])
add_text(slide, "资金使用原则", 4.72, 2.12, 1.8, 0.28, size=21, color=COLORS["navy"], bold=True, font=TITLE_FONT)
add_bullets(slide, ["首批只下核心 SKU，不多囤货", "先打样和小批验证，再决定第二批", "第一批回款优先反哺第二批下单", "生活开销压到最低，保证项目续航"], 4.68, 2.55, 3.4, 2.1, size=14)
add_text(slide, "第 8 周是最大现金支出节点，建议分两批下单，避免一次性掏空。", 4.72, 5.2, 3.25, 0.46, size=12.5, color=COLORS["ink"])
add_round_rect(slide, 8.78, 1.9, 3.55, 4.55, COLORS["cream"], line=COLORS["cream"])
add_text(slide, "阶段性支出曲线", 9.02, 2.12, 2.1, 0.28, size=21, color=COLORS["navy"], bold=True, font=TITLE_FONT)
for i, (wk, val) in enumerate([("第4周", "累计 ¥4,350"), ("第8周", "累计 ¥22,050"), ("第10周", "累计 ¥23,550"), ("第12周", "累计 ¥26,050")]):
    y = 2.66 + i * 0.8
    add_circle(slide, 9.08, y, 0.3, COLORS["coral"])
    add_text(slide, wk, 9.48, y - 0.02, 0.9, 0.22, size=13, color=COLORS["navy"], bold=True)
    add_text(slide, val, 10.28, y - 0.02, 1.45, 0.22, size=13, color=COLORS["ink"])
add_text(slide, "目标：第 12 周仍保留约 ¥3,950 运营资金。", 9.02, 5.45, 2.8, 0.3, size=12.5, color=COLORS["muted"])
add_footer(slide, "来源：06_启动资金使用明细.md", dark=True)

# Slide 9
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, COLORS["cream"])
add_text(slide, "08｜制作落地：先验证可做，再验证稳定", 0.72, 0.58, 7.8, 0.6, size=28, color=COLORS["navy"], bold=True, font=TITLE_FONT)
add_text(slide, "生产方案把流程收敛为五个节点，每个节点都有明确输出物和闸口。", 0.72, 1.1, 7.5, 0.36, size=15, color=COLORS["muted"])
phases = [
    ("打样", "7–14天", "样品 / 报价 / 修订意见", COLORS["rose"]),
    ("试产", "7–10天", "试产批 / 良率数据", COLORS["gold"]),
    ("大货", "10–21天", "首批可售库存", COLORS["sand"]),
    ("质检", "全程并行", "进料检 / 在线检 / 成品检", COLORS["rose"]),
    ("入库", "1–3天", "入库清单 / 上架准备", COLORS["gold"]),
]
start_x = 0.95
for i, (title, days, out, fill) in enumerate(phases):
    x = start_x + i * 2.45
    add_circle(slide, x + 0.75, 2.0, 0.72, COLORS["navy"], text=str(i + 1), text_size=18)
    add_round_rect(slide, x, 2.88, 1.95, 2.02, fill, line=fill)
    add_text(slide, title, x + 0.18, 3.08, 1.59, 0.28, size=18, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, days, x + 0.18, 3.5, 1.59, 0.26, size=14, color=COLORS["coral"], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, out, x + 0.16, 3.95, 1.63, 0.64, size=11.5, color=COLORS["ink"], align=PP_ALIGN.CENTER)
    if i < len(phases) - 1:
        conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x + 1.95), Inches(3.88), Inches(x + 2.42), Inches(3.88))
        conn.line.color.rgb = COLORS["line"]
        conn.line.width = Pt(2.2)
add_round_rect(slide, 0.95, 5.52, 11.9, 0.74, COLORS["white"], line=COLORS["line"])
add_text(slide, "三条硬闸口：样品不通过不进试产；试产良率不达标不进大货；包装与大货未确认不进渠道。", 1.2, 5.77, 11.35, 0.2, size=14, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)
add_footer(slide, "来源：设计/03_执行手册/09_制作落地执行方案.md")

# Slide 10
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, COLORS["cream"])
add_text(slide, "09｜12 周执行节奏：每周都要有产出与复盘", 0.72, 0.58, 8.8, 0.6, size=28, color=COLORS["navy"], bold=True, font=TITLE_FONT)
add_text(slide, "执行表的价值不在于排满，而在于让每周优先解决影响下周排期的关键节点。", 0.72, 1.1, 9.6, 0.36, size=15, color=COLORS["muted"])
for x, fill, title, bullets in [
    (0.92, COLORS["rose"], "第1–2周｜样稿启动", ["确定洪洞首轮样板方向", "完成纸品布艺版式初稿与供应商询价", "同步准备小红书账号与测试内容"]),
    (4.42, COLORS["gold"], "第3–4周｜小批验证", ["根据反馈微调版式", "线下试卖 / 小范围试看", "形成继续 / 回退 / 暂停结论"]),
    (7.92, COLORS["sand"], "第5–12周｜按结论推进", ["补足样稿、优化供应商与成本结构", "决定是否复制到更多产品目录", "建立标准化执行模板"]),
]:
    add_round_rect(slide, x, 1.88, 3.0, 3.95, fill, line=fill)
    add_text(slide, title, x + 0.18, 2.1, 2.5, 0.45, size=19, color=COLORS["navy"], bold=True, font=TITLE_FONT)
    add_bullets(slide, bullets, x + 0.18, 2.75, 2.45, 1.8, size=13.5)
add_round_rect(slide, 0.95, 6.05, 11.35, 0.48, COLORS["navy"], line=COLORS["navy"])
add_text(slide, "每周日晚上必须完成三件事：记账、总结当周、规划下周。", 1.15, 6.16, 10.95, 0.16, size=13.5, color=COLORS["cream"], bold=True, align=PP_ALIGN.CENTER)
add_footer(slide, "来源：07_每周执行清单.md")

# Slide 11
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, COLORS["cream"])
add_text(slide, "10｜渠道与合作：先试点，再放量", 0.72, 0.58, 7.5, 0.6, size=28, color=COLORS["navy"], bold=True, font=TITLE_FONT)
add_text(slide, "不同品类适配不同成交场景，但共用同一推进原则：先样品，再试销，再复盘。", 0.72, 1.1, 8.6, 0.36, size=15, color=COLORS["muted"])
channels = [
    ("景区门店", "低客单引流 / DIY体验 / 现场试戴", COLORS["rose"]),
    ("私域社群", "故事内容种草 / 预约定制 / 复购", COLORS["gold"]),
    ("电商渠道", "标准化礼盒 / 套装 / 可补货产品", COLORS["sand"]),
    ("合作门店", "唐装联名 / 陈列样品 / 高客单试点", COLORS["rose"]),
]
for i, (title, desc, fill) in enumerate(channels):
    x = 0.95 + (i % 2) * 6.0
    y = 1.9 + (i // 2) * 1.95
    add_round_rect(slide, x, y, 5.2, 1.45, fill, line=fill)
    add_circle(slide, x + 0.2, y + 0.26, 0.46, COLORS["navy"])
    add_text(slide, title, x + 0.82, y + 0.2, 1.6, 0.28, size=18, color=COLORS["navy"], bold=True)
    add_text(slide, desc, x + 0.82, y + 0.6, 3.95, 0.4, size=13, color=COLORS["ink"])
add_text(slide, "推进共识", 1.02, 6.1, 1.4, 0.25, size=14, color=COLORS["coral"], bold=True)
add_bullets(slide, [
    "门店反馈弱：先改配搭与话术，不急于扩 SKU",
    "渠道验证通过：再决定补货、礼盒化与联名放量",
    "高客单线优先做样品与体验感，不直接压大货",
], 1.7, 6.0, 10.6, 0.8, size=13)
add_footer(slide, "来源：唐装合作、DIY体验计划、业务逻辑总览")

# Slide 12
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, COLORS["navy"])
add_text(slide, "11｜建议的下一步动作", 0.78, 0.7, 4.6, 0.58, size=30, color=COLORS["white"], bold=True, font=TITLE_FONT)
add_text(slide, "用 30 天完成一个可验证闭环：选品 → 打样 → 试销 → 复盘。", 0.78, 1.28, 5.8, 0.34, size=16, color=COLORS["cream"])
next_steps = [
    ("收窄首发 SKU", "地域线先做纸品布艺 4 款；手工线先做 3 类样板"),
    ("并行准备样品", "同步完成 DIY 三款主打和唐装门店试戴样品"),
    ("建立试销记录", "记录成本、工时、良率、成交反馈与内容传播效果"),
    ("通过闸口再放量", "样品和小批验证通过后，再扩 SKU、礼盒和联名"),
]
for i, (title, desc) in enumerate(next_steps):
    y = 2.0 + i * 1.12
    add_circle(slide, 0.95, y, 0.48, COLORS["gold"], text=str(i + 1), text_size=16, text_color=COLORS["navy"])
    add_text(slide, title, 1.6, y - 0.02, 2.2, 0.24, size=18, color=COLORS["gold"], bold=True)
    add_text(slide, desc, 3.75, y - 0.01, 4.5, 0.28, size=13.5, color=COLORS["cream"])
    if i < len(next_steps) - 1:
        conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.19), Inches(y + 0.48), Inches(1.19), Inches(y + 0.95))
        conn.line.color.rgb = COLORS["sage"]
        conn.line.width = Pt(2)
add_round_rect(slide, 9.4, 1.75, 3.0, 3.85, COLORS["cream"], line=COLORS["cream"])
add_text(slide, "结论", 9.8, 2.18, 1.2, 0.28, size=22, color=COLORS["navy"], bold=True, font=TITLE_FONT)
add_text(slide, "这套资料已经不是零散笔记，而是一套可直接进入选品、打样和合作执行的项目底稿。", 9.8, 2.82, 2.15, 1.18, size=16, color=COLORS["ink"])
add_round_rect(slide, 9.8, 4.42, 1.8, 0.5, COLORS["gold"], line=COLORS["gold"])
add_text(slide, "先验证，再放量", 9.98, 4.53, 1.44, 0.16, size=13, color=COLORS["navy"], bold=True, align=PP_ALIGN.CENTER)
add_footer(slide, "总结依据：产品索引、计划索引、执行手册与专题计划", dark=True)

prs.save(OUT)
print(OUT)
